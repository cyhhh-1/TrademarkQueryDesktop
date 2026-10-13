"""Verify generated XLSX deliverables without relying on Excel automation."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
import sys
import zipfile

from openpyxl import load_workbook


def inspect(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        external_relationships = [
            name
            for name in archive.namelist()
            if name.endswith(".rels")
            and b'TargetMode="External"' in archive.read(name)
        ]
        media = [
            name for name in archive.namelist() if name.startswith("xl/media/")
        ]

    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheet = workbook.active
        image = sheet._images[0]
        anchor = image.anchor
        last_column = sheet.cell(1, sheet.max_column).column_letter
        return {
            "path": str(path),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "bad_member": bad_member,
            "values": [
                sheet.cell(2, column).value
                for column in range(1, sheet.max_column + 1)
            ],
            "image_count": len(sheet._images),
            "anchor": type(anchor).__name__,
            "editAs": getattr(anchor, "editAs", None),
            "from": [anchor._from.row, anchor._from.col],
            "to": [anchor.to.row, anchor.to.col],
            "row_height": sheet.row_dimensions[2].height,
            "column_width": sheet.column_dimensions[last_column].width,
            "image_size": [image.width, image.height],
            "external_relationships": external_relationships,
            "media": media,
        }
    finally:
        workbook.close()


def main() -> None:
    paths = [Path(argument) for argument in sys.argv[1:]]
    if not paths:
        raise SystemExit("usage: verify-deliverables.py OUTPUT.xlsx [...]")
    print(json.dumps([inspect(path) for path in paths], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
