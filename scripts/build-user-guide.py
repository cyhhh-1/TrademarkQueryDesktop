"""Build the employee and administrator guide for Trademark Query Desktop."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BLUE = "145DA0"
PALE_BLUE = "EEF5FC"
PALE_GRAY = "F6F7F9"
BORDER = "D9D9D9"
TEXT = RGBColor(0, 0, 0)
MUTED = RGBColor(80, 80, 80)


def set_run_font(
    run,
    size: float | None = None,
    bold: bool | None = None,
    east_asia: str = "宋体",
) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.color.rgb = TEXT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 110, start: int = 130, bottom: int = 110, end: int = 130) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def keep_table_row_together(row) -> None:
    properties = row._tr.get_or_add_trPr()
    if properties.find(qn("w:cantSplit")) is None:
        properties.append(OxmlElement("w:cantSplit"))


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, 10.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, value, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 10.5)


def add_paragraph(document, text: str = "", *, bold_lead: str | None = None, style: str | None = None):
    paragraph = document.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        run = paragraph.add_run(text[len(bold_lead):])
        set_run_font(run)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullets(document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run)


def add_numbered_steps(document, items: list[tuple[str, str]]) -> None:
    for index, (title, detail) in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.62)
        paragraph.paragraph_format.first_line_indent = Cm(-0.62)
        paragraph.paragraph_format.space_after = Pt(5)
        number_run = paragraph.add_run(f"{index}. ")
        set_run_font(number_run, 10.5)
        title_run = paragraph.add_run(title)
        set_run_font(title_run, 10.5, True)
        detail_run = paragraph.add_run(f"  {detail}")
        set_run_font(detail_run, 10.5)


def add_table(
    document,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
):
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, 10.5)
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, label) in enumerate(zip(header.cells, headers)):
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade_cell(cell, BLUE)
        set_cell_borders(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        set_run_font(run, 10.5, True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        keep_table_row_together(row)
        cells = row.cells
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade_cell(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            set_cell_borders(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            set_run_font(run, 10.5)
    document.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(
    document,
    path: Path,
    caption: str,
    alt_text: str,
    width_inches: float = 6.0,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width_inches))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    paragraph.paragraph_format.keep_with_next = True
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(8)
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, 10.5)
    caption_run.italic = True
    caption_run.font.color.rgb = MUTED


def add_code_block(document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.6)
    shade_cell(cell, PALE_GRAY)
    set_cell_borders(cell, "E3E6EA", "4")
    set_cell_margins(cell, 120, 150, 120, 150)
    paragraph = cell.paragraphs[0]
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8.5)
        run.font.color.rgb = TEXT
        if index < len(lines) - 1:
            run.add_break()
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    title = styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_properties = title._element.get_or_add_pPr()
    title_borders = title_properties.find(qn("w:pBdr"))
    if title_borders is not None:
        title_properties.remove(title_borders)

    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(6)

    for name in ("List Number", "List Bullet"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12)
        style.paragraph_format.space_after = Pt(0)

    if "Small Meta" not in styles:
        meta = styles.add_style("Small Meta", WD_STYLE_TYPE.PARAGRAPH)
        meta.font.name = "Times New Roman"
        meta._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        meta.font.size = Pt(10.5)
        meta.font.color.rgb = MUTED

    header = section.header
    header.distance = Cm(0.7)
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = "商标查询助手 0.1.20  安装与使用说明"
    for run in header_paragraph.runs:
        set_run_font(run, 10.5)
        run.font.color.rgb = MUTED
    footer = section.footer
    footer.distance = Cm(0.65)
    add_page_number(footer.paragraphs[0])


def build(output: Path) -> None:
    document = Document()
    configure_document(document)
    properties = document.core_properties
    properties.title = "商标查询助手安装与使用说明"
    properties.subject = "员工使用和管理员分发手册"
    properties.author = "商标查询助手"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("商标查询助手安装与使用说明")
    set_run_font(title_run, 22, True, "黑体")
    subtitle = document.add_paragraph(style="Small Meta")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("适用版本 0.1.20    更新日期 2026年9月28日")
    opening = add_paragraph(
        document,
        "本手册面向使用商标查询助手的员工和负责分发安装包的管理员。程序完全在员工电脑本地运行，不需要启动 BuddyAgent 后端、管理端前端、MQTT、技能服务或数据库。每位员工使用自己的国家知识产权局账号登录，查询结果由员工自行另存。",
    )
    opening.paragraph_format.space_before = Pt(8)
    opening.paragraph_format.space_after = Pt(10)

    document.add_heading("1 快速开始", level=1)
    add_numbered_steps(
        document,
        [
            ("安装并打开", "运行安装程序，安装完成后打开“商标查询助手”。"),
            ("选择表格", "点击“选择 XLSX”，选择商标申请或商标唯一性查询表格。"),
            ("打开网站", "点击“打开商标网站”，程序会打开一份仅供本应用使用的 Microsoft Edge。"),
            ("员工自行登录", "在专用 Edge 中点击“商标综合查询”，自行输入账号、密码和验证码；登录成功后网站会自动返回商标首页。"),
            ("检查并开始", "回到助手，点击“我已登录，检查并进入查询”；状态显示可开始后，点击“开始查询”。"),
            ("观察并保存", "通过当前商标、进度、阶段、耗时和预计剩余时间确认任务仍在运行；完成后点击“另存为”。"),
        ],
    )
    add_paragraph(
        document,
        "重要：开始查询后，专用 Edge 会移到屏幕外并在后台工作。可以正常使用其他软件和自己的浏览器，但不要关闭商标查询助手或专用 Edge。单条异常会在末尾自动重试，不会暂停整批任务；任务结束后 Windows 会显示消息通知。",
        bold_lead="重要：",
    )
    document.add_page_break()

    document.add_heading("2 使用前准备", level=1)
    document.add_heading("电脑和账号", level=2)
    add_bullets(
        document,
        [
            "Windows 10 或 Windows 11，64 位系统。",
            "已安装 Microsoft Edge。程序使用员工电脑上的普通 Edge，不使用内置浏览器。",
            "建议安装 Microsoft Excel；WPS 可用于查看结果，但首次验收建议用 Excel。",
            "每位员工使用自己的国家知识产权局账号。多人可分别在各自电脑同时运行。",
            "网络能正常访问 https://wcjs.sbj.cnipa.gov.cn/ 和统一身份认证页面。",
        ],
    )
    document.add_heading("两种输入表格", level=2)
    add_table(
        document,
        "表 1  输入表格与输出内容",
        ["任务", "输入要求", "程序写入结果"],
        [
            [
                "商标申请查询",
                "每行填写商标名称和商标申请人。中文申请人自动填入网站中文栏，英文申请人自动填入英文栏。",
                "申请人、专用权期限、商标详情截图。",
            ],
            [
                "商标唯一性查询",
                "每行填写商标名称；保留是否唯一性、商标申请人、商标有效期和截图表头。",
                "是否唯一、同一公司时的申请人和有效期、全部同名商标行的纵向拼接截图。",
            ],
        ],
        [1.1, 2.75, 2.15],
    )
    add_paragraph(
        document,
        "一份表格可以包含 1 个、多个或十几个以上商标。程序按工作表顺序逐行处理，每完成一行立即写入本地工作副本并保存检查点。请勿改名、删除或合并模板表头，也不要在查询期间用 Excel 打开结果工作副本。",
    )

    document.add_heading("3 安装和首次打开", level=1)
    add_numbered_steps(
        document,
        [
            ("解压发布包", "把管理员提供的 ZIP 完整解压到普通文件夹，不要直接在压缩包预览中运行安装程序。"),
            ("核对版本", "确认安装文件名为“商标查询助手_0.1.20_x64-setup.exe”，并按管理员提供的 SHA256 文件核验。"),
            ("运行安装", "双击安装程序。该版本按当前 Windows 用户安装，不需要启动公司现有项目。"),
            ("安全提示", "若 Windows 或安全软件阻止运行，不要自行绕过；把提示截图发给管理员，由管理员按公司软件策略签名、放行或重新分发。"),
            ("打开应用", "安装完成后从开始菜单或桌面入口打开“商标查询助手”。"),
        ],
    )
    document.add_heading("4 执行批量查询", level=1)
    document.add_heading("选择 XLSX", level=2)
    add_numbered_steps(
        document,
        [
            ("点击选择 XLSX", "程序自动识别两种表格模板，并列出所有非空商标行。"),
            ("核对数量、商标和申请人", "查看左侧文件名、检测到的模板和总行数；在右侧列表滚动核对商标及其申请人。唯一性查询没有输入申请人时显示“-”。"),
            ("纠正输入问题", "若提示缺少表头、申请人或文件损坏，先关闭助手，修正原表后重新选择。"),
        ],
    )
    document.add_heading("登录国家知识产权局网站", level=2)
    add_numbered_steps(
        document,
        [
            ("打开商标网站", "助手打开专用 Edge 和商标网站首页，并自动最小化自身让登录窗口显示在前面。不要改用日常 Edge，因为登录会话不共享。"),
            ("进入综合查询", "在网站首页点击“商标综合查询”。需要登录时会进入统一身份认证页面。"),
            ("员工自行完成登录", "输入本人账号、密码和图片验证码后点击登录。账号、密码和验证码不经过本应用。"),
            ("等待自动返回", "登录成功后等待网站自动返回商标首页。若变成 about:blank 或 403，回到首页重新进入，必要时稍后再试。"),
            ("回到助手确认", "点击“我已登录，检查并进入查询”。助手会通过商标网站首页进入 /sgtmi 查询页；状态显示“可以开始”后再点击“开始查询”。"),
        ],
    )

    document.add_page_break()
    document.add_heading("5 看懂进度和预计时间", level=1)
    add_table(
        document,
        "表 2  查询进度信息说明",
        ["界面信息", "含义"],
        [
            ["当前处理", "正在查询的商标名称；对应列表行会蓝色高亮并自动滚动到可见位置。"],
            ["商标申请人", "显示输入表格中与商标同一行的申请人；唯一性查询确认唯一申请人后会更新。"],
            ["第几项 共多少项", "本次任务的整体位置。暂停后继续仍保持原来的总数和序号。"],
            ["完成进度", "按已经安全写入结果表格的行数计算。当前行尚未写入时百分比不会提前增加。"],
            ["当前阶段", "包括检索、核对结果、打开详情、生成证据截图和写入结果表格。"],
            ["已用时间", "显示程序实际工作时间；员工主动暂停的时间不计入。"],
            ["预计剩余", "完成第一个商标后，根据已完成平均耗时动态估算；网站速度变化时会调整。"],
        ],
        [1.45, 4.55],
    )
    add_paragraph(
        document,
        "单个商标可能需要打开多条同名记录和详情页，唯一性查询还会遍历结果页并裁剪全部同名行，因此不同商标耗时可能差异很大。只要当前商标、阶段或已用时间仍在更新，就说明任务仍在执行。",
    )
    add_paragraph(
        document,
        "查询运行时，选择新文件、移除文件和重新打开网站按钮会锁定，防止误操作。专用 Edge 会移到屏幕外并在后台工作；员工可以最小化助手并使用其他应用或日常浏览器，但不要手工操作这一个专用 Edge。点击“取消任务”后，程序会先完成当前商标并安全写入，再停止后续行；已经完成的阶段性结果可另存。",
    )

    document.add_heading("6 输出与结果验证", level=1)
    document.add_heading("商标申请查询", level=2)
    add_bullets(
        document,
        [
            "检索要素填写商标名称，并填写表格中的申请人名称。",
            "只选择商标名称和申请人都一致的结果。",
            "同名结果有多条时，优先从序号较大的记录开始检查；没有专用权期限时继续检查较早记录。",
            "找到申请人一致且含专用权期限的详情后，填写申请人和专用权期限，并把完整详情页截图嵌入该行截图列。",
        ],
    )
    document.add_heading("商标唯一性查询", level=2)
    add_bullets(
        document,
        [
            "仅按商标名称检索，并遍历所有结果页。",
            "只保留与检索词标准化后完全相同的商标，逐条裁剪并纵向拼接证据。",
            "全部同名商标的申请人标准化后只有一家时填写“是”；存在多家公司时填写“否”。",
            "填写“是”时再从较大序号开始检查详情，补充申请人和有效期；填写“否”时申请人和有效期保持空白。",
        ],
    )
    document.add_heading("截图在表格中的表现", level=2)
    add_paragraph(
        document,
        "截图文件保存在 XLSX 内部，不依赖电脑上的外部图片路径。图片采用随单元格移动和调整大小的嵌入方式，Excel 和 WPS 打开、复制或移动结果文件时仍能显示。程序会保持原图比例，必要时补白边，并把单行高度限制在 Excel 支持的 409 点以内；唯一性结果很长时，XLSX 内仍保留原始分辨率，放大图片或提取图片后文字不会因插入过程被永久压糊。",
    )
    add_paragraph(
        document,
        "结果会先写入应用本地工作目录。任务完成或安全取消后点击“另存为”，选择员工自己的业务文件夹。程序不会覆盖原始输入表格。",
    )

    document.add_page_break()
    document.add_heading("7 暂停 恢复和常见问题", level=1)
    add_table(
        document,
        "表 3  常见问题与处理方法",
        ["现象", "处理方法"],
        [
            ["任务开始前提示需要登录", "点击“打开商标网站”，在专用 Edge 重新登录；回到助手点击“我已登录，检查并进入查询”。"],
            ["网站已打开却显示“后台截图服务未能启动”", "关闭旧版助手，安装 0.1.20 后重新点击“打开商标网站”；新版会接回同一专用 Edge 的登录会话，不必清除浏览器资料。"],
            ["已经登录却仍提示先打开网站", "确认使用 0.1.20 或更高版本；登录后点击“我已登录”，助手会从商标网站首页进入综合查询。仍未恢复时检查是否出现验证码、拒绝访问或重新登录页面。"],
            ["任务中个别商标出现错误", "无需处理。程序会继续查询后续商标，并在首轮结束后自动重试失败项；最终失败或未找到时，具体原因会写入原表的截图列，并同时记录在“查询执行记录”中。"],
            ["查询时还要处理其他工作", "可以最小化助手并正常使用其他应用或日常浏览器。不要在专用 Edge 中新开网页、切换标签、关闭标签或修改地址；任务开始后助手会锁定重新打开网站按钮。"],
            ["重复打开了两个客户端", "程序只允许一个 Worker 使用专用登录会话；后打开的客户端会提示已有客户端正在运行。关闭重复窗口，回到先打开的客户端继续。"],
            ["出现图片验证码", "若查询中遇到验证码，该商标会记录失败并继续处理其他行；任务结束后由员工打开专用 Edge 人工处理并复查。助手不会识别、提交或保存验证码。"],
            ["出现 403 或请求被阻断", "程序会等待后自动刷新结果页一次；刷新后仍被阻断时先继续后续商标，并在首轮结束后完整重试一次。完整重试仍失败时，错误原因会写入截图列。"],
            ["页面变成 about:blank", "返回商标网站首页重新进入综合查询；完成登录后确认页面确实回到首页，再回助手检查。"],
            ["预计剩余时间未显示", "这是正常现象。第一个商标安全写入后才开始估算。"],
            ["某一商标停留较久", "查看当前阶段。多结果、详情页加载、网站限速和证据拼接都会增加耗时。"],
            ["无法另存或文件被占用", "关闭 Excel 或 WPS 中同名结果文件，再选择新的保存路径。不要覆盖正在打开的文件。"],
            ["个别商标没有结果", "核对商标和申请人是否完全正确；唯一性任务还要确认表头完整。必要时人工复核网站。"],
            ["应用或 Edge 被意外关闭", "不要继续操作原输入表。重新打开应用和表格重新开始；本地任务目录仍保留已写入的工作副本供管理员排查。"],
        ],
        [1.65, 4.35],
    )

    document.add_heading("8 管理员分发和配置", level=1)
    document.add_heading("建议的分发方式", level=2)
    add_numbered_steps(
        document,
        [
            ("使用发布包", "把 ZIP 放到公司网盘、飞书云盘或受控共享目录；不要让员工从源码目录运行。"),
            ("保留完整文件", "发布包应包含安装程序、Word 手册、PDF 手册和 SHA256 校验文件。"),
            ("先做一台试装", "使用普通员工账号安装，分别加载两种模板，验证 Edge 登录、查询和另存。"),
            ("再批量分发", "每位员工在自己的电脑安装并在需要时启动，不设置随系统自动启动。"),
            ("升级版本", "后续发布新版时使用新的完整发布包，并保留版本号和校验文件，避免混用旧安装程序。"),
        ],
    )
    document.add_heading("客户端中的大模型设置", level=2)
    add_paragraph(
        document,
        "0.1.20 安装包已经预置公司统一的大模型接口、API Key 和模型名称。员工安装后可以直接查询，无需编辑配置文件。该配置只在网页中的商标文字需要图片识别时作为兜底。",
    )
    add_numbered_steps(
        document,
        [
            ("打开设置", "点击客户端右上角的“模型设置”。"),
            ("查看或修改", "接口 URL、API Key 和模型名称都可以直接输入；API Key 默认遮罩，点击“显示”后才会出现明文。"),
            ("保存配置", "确认内容后点击“保存配置”，新配置从下一个查询任务开始使用。"),
            ("恢复默认", "填写错误或需要恢复统一配置时，点击“恢复公司默认”，再点击“保存配置”。"),
        ],
    )
    add_paragraph(
        document,
        "安全要求：公司 Key 会随内部安装包分发，客户端会遮罩显示，但无法对拥有本机权限的安装包使用者绝对保密。发布包只应提供给内部员工；密钥轮换后应通过客户端逐台修改，或重新发布带新默认配置的安装包。不要在截图、聊天记录或公开资料中发送真实 Key。",
        bold_lead="安全要求：",
    )
    document.add_heading("与公司现有项目的关系", level=2)
    add_bullets(
        document,
        [
            "这是独立桌面项目和独立安装包，不修改 BuddyAgent 源码。",
            "员工使用时不依赖 ai-agent-admin、后端服务、数据库、MQTT、技能或 SOP。",
            "每台电脑拥有独立 Edge 配置、登录会话、任务目录和设置文件；多人并发互不占用同一浏览器。",
            "安装包内置公司千问配置值，但运行时不会读取原项目文件。",
        ],
    )

    document.add_heading("9 本地数据 隐私和支持", level=1)
    add_table(
        document,
        "表 4  本地数据位置与处理方式",
        ["内容", "本地位置或处理方式"],
        [
            ["Edge 登录资料", "%LOCALAPPDATA%\\TrademarkQueryDesktop\\edge-profile"],
            ["任务检查点 临时截图 工作副本", "%LOCALAPPDATA%\\TrademarkQueryDesktop\\tasks"],
            ["模型设置", "%LOCALAPPDATA%\\TrademarkQueryDesktop\\settings.json；普通员工通过客户端界面修改。"],
            ["账号 密码 验证码", "由 Edge 和国家知识产权局网站管理；本应用不读取或保存。"],
            ["最终结果", "由员工点击“另存为”保存到自己选择的位置。"],
        ],
        [2.0, 4.0],
    )
    add_paragraph(
        document,
        "任务完成并确认已另存后，管理员可按公司数据保留策略清理本地 tasks 目录。清理前先关闭商标查询助手和专用 Edge，并确认不再需要检查点、证据和工作副本。",
    )
    document.add_heading("报障时提供的信息", level=2)
    add_bullets(
        document,
        [
            "应用版本、发生时间和 Windows 版本。",
            "界面完整截图，以及专用 Edge 当前页面截图。",
            "发生问题的商标名称、XLSX 行号和当前阶段。",
            "是否能从商标网站首页手工进入综合查询。",
            "不要发送账号密码、验证码、真实 API 密钥或无关个人信息。",
        ],
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
