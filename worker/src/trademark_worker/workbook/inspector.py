"""Detect supported XLSX templates and construct row-level jobs."""

from __future__ import annotations

import re
import shutil
import stat
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from trademark_worker.models import LookupJob, LookupMode


class WorkbookInputError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class WorkbookInspection:
    mode: LookupMode
    worksheet: str
    header_row: int
    columns: dict[str, int]
    task_count: int

    def to_wire(self) -> dict[str, object]:
        return {
            "mode": self.mode.value,
            "worksheet": self.worksheet,
            "header_row": self.header_row,
            "columns": self.columns,
            "task_count": self.task_count,
        }


def normalize_header(value: object) -> str:
    normalized = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"[\s()（）\[\]【】“”\"']+", "", normalized).casefold()


def _column_matching(headers: list[str], prefixes: tuple[str, ...]) -> int | None:
    for index, header in enumerate(headers, start=1):
        if any(header == prefix or header.startswith(prefix) for prefix in prefixes):
            return index
    return None


def _require_columns(
    headers: list[str], definitions: dict[str, tuple[str, tuple[str, ...]]]
) -> dict[str, int]:
    columns: dict[str, int] = {}
    missing: list[str] = []
    for key, (label, prefixes) in definitions.items():
        found = _column_matching(headers, prefixes)
        if found is None:
            missing.append(label)
        else:
            columns[key] = found
    if missing:
        raise WorkbookInputError(
            "TEMPLATE_MISSING_HEADER", "缺少表头：" + "、".join(missing)
        )
    return columns


def _find_header_sheet(workbook) -> tuple[object, list[str]]:
    for sheet in workbook.worksheets:
        headers = [normalize_header(cell.value) for cell in sheet[1]]
        if any(headers):
            return sheet, headers
    raise WorkbookInputError("TEMPLATE_EMPTY", "未找到包含表头的工作表")


def inspect_workbook(path: Path) -> WorkbookInspection:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise WorkbookInputError("UNSUPPORTED_FILE_TYPE", "只支持 .xlsx 文件")
    if not path.is_file():
        raise WorkbookInputError("FILE_NOT_FOUND", "找不到 XLSX 文件")
    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:
        raise WorkbookInputError("WORKBOOK_DAMAGED", "XLSX 文件无法打开") from exc
    try:
        sheet, headers = _find_header_sheet(workbook)
        is_uniqueness = any(header.startswith("是否唯一性") for header in headers)
        if is_uniqueness:
            mode = LookupMode.UNIQUENESS
            columns = _require_columns(
                headers,
                {
                    "mark": ("商标名称", ("商标名称",)),
                    "uniqueness": ("是否唯一性", ("是否唯一性",)),
                    "applicant": ("商标申请人", ("商标申请人",)),
                    "term": ("商标有效期", ("商标有效期", "专用权期限")),
                    "screenshot": ("截图", ("截图",)),
                },
            )
        else:
            mode = LookupMode.APPLICATION
            columns = _require_columns(
                headers,
                {
                    "mark": ("商标名称", ("商标名称",)),
                    "applicant": ("商标申请人", ("商标申请人",)),
                    "term": ("专用权期限", ("专用权期限", "商标有效期")),
                    "screenshot": ("截图", ("截图",)),
                },
            )
        task_count = sum(
            1
            for row in range(2, sheet.max_row + 1)
            if _cell_text(sheet.cell(row, columns["mark"]).value)
        )
        return WorkbookInspection(mode, sheet.title, 1, columns, task_count)
    finally:
        workbook.close()


def prepare_working_copy(source: Path, run_dir: Path) -> Path:
    source = Path(source)
    if not source.is_file():
        raise WorkbookInputError("FILE_NOT_FOUND", "找不到 XLSX 文件")
    run_dir = Path(run_dir)
    input_dir = run_dir / "input"
    output_dir = run_dir / "output"
    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)
    snapshot = input_dir / source.name
    if not snapshot.exists():
        shutil.copy2(source, snapshot)
    output_stem = source.stem if source.stem.endswith("_已查询") else f"{source.stem}_已查询"
    working = output_dir / f"{output_stem}.xlsx"
    if not working.exists():
        shutil.copy2(snapshot, working)
    # copy2 preserves the source's Windows read-only attribute. The input and
    # snapshot stay untouched, but the output must permit atomic replacement.
    working.chmod(working.stat().st_mode | stat.S_IWRITE)
    return working


def load_jobs(
    path: Path, inspection: WorkbookInspection, overwrite: bool = False
) -> list[LookupJob]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        sheet = workbook[inspection.worksheet]
        jobs: list[LookupJob] = []
        for row in range(inspection.header_row + 1, sheet.max_row + 1):
            mark = _cell_text(sheet.cell(row, inspection.columns["mark"]).value)
            if not mark:
                continue
            result_key = "uniqueness" if inspection.mode is LookupMode.UNIQUENESS else "term"
            has_result = _cell_text(
                sheet.cell(row, inspection.columns[result_key]).value
            )
            if has_result and not overwrite:
                continue
            if inspection.mode is LookupMode.APPLICATION:
                applicant = _cell_text(
                    sheet.cell(row, inspection.columns["applicant"]).value
                )
                if not applicant:
                    raise WorkbookInputError(
                        "ROW_MISSING_APPLICANT", f"第 {row} 行缺少商标申请人"
                    )
                jobs.append(
                    LookupJob.application(row, mark, applicant, inspection.worksheet)
                )
            else:
                jobs.append(LookupJob.uniqueness(row, mark, inspection.worksheet))
        return jobs
    finally:
        workbook.close()


def _cell_text(value: object) -> str:
    return str(value).strip() if value is not None else ""
