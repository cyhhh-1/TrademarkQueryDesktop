from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from threading import Event, Thread

from openpyxl import Workbook, load_workbook
from PIL import Image

from trademark_worker.checkpoint import CheckpointStore
from trademark_worker.models import (
    BrowserState,
    DetailRecord,
    LookupJob,
    LookupMode,
    NeedsUserAction,
    ResultRecord,
    TaskRequest,
    TaskState,
)
from trademark_worker.task_engine import TaskEngine


def make_workbook(tmp_path: Path, mode: LookupMode, rows: list[list[object]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    if mode is LookupMode.APPLICATION:
        sheet.append(["商标名称", "商标申请人", "专用权期限", "截图"])
        name = "商标申请.xlsx"
    else:
        sheet.append(
            [
                "商标名称",
                "是否唯一性\n（当申请公司有多家时填写“否”",
                "商标申请人\n当是否唯一性“否”时此处空白",
                "商标有效期",
                "截图",
            ]
        )
        name = "商标查询是否唯一性.xlsx"
    for row in rows:
        sheet.append(row)
    path = tmp_path / name
    workbook.save(path)
    return path


def record(
    sequence: int,
    mark: str,
    applicant: str,
    registration_no: str | None = None,
) -> ResultRecord:
    return ResultRecord(
        sequence,
        registration_no or f"9000{sequence}",
        "35",
        "2026年08月03日",
        mark,
        applicant,
        page_number=1,
        row_index=sequence - 1,
    )


def detail(source: ResultRecord, term: str | None) -> DetailRecord:
    return DetailRecord(
        source.mark,
        source.registration_no,
        source.applicant,
        term,
    )


def active_detail(source: ResultRecord, term: str = "2025年至2035年") -> DetailRecord:
    return DetailRecord(
        source.mark,
        source.registration_no,
        source.applicant,
        term,
        status="LIVE/REGISTRATION/Issued and Active 注册商标",
    )


class FakeClient:
    def __init__(
        self,
        records: dict[str, list[ResultRecord]],
        details: dict[int, DetailRecord] | None = None,
        pause_on_mark: str | None = None,
    ) -> None:
        self.records = records
        self.details = details or {}
        self.pause_on_mark = pause_on_mark
        self.opened_sequences: list[int] = []
        self.stacked_sequences: list[int] = []

    def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
        if job.mark == self.pause_on_mark:
            raise NeedsUserAction(BrowserState.LOGIN_REQUIRED, "请登录")
        return self.records[job.mark]

    def search_pages(self, job: LookupJob, page_progress=None):
        yield self.search(job, page_progress)

    def open_detail(self, source: ResultRecord) -> DetailRecord:
        self.opened_sequences.append(source.sequence)
        return self.details[source.sequence]

    def capture_exact_rows(
        self, records: list[ResultRecord], output_dir: Path
    ) -> Path:
        self.stacked_sequences.extend(item.sequence for item in records)
        path = output_dir / "exact-results.png"
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (600, 80 * len(records)), "white").save(path)
        return path

    def capture_detail(self, output: Path) -> Path:
        output.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (600, 800), "white").save(output)
        return output


class PaginatedFakeClient(FakeClient):
    def __init__(
        self,
        pages: dict[str, list[list[ResultRecord]]],
        details: dict[int, DetailRecord],
    ) -> None:
        super().__init__(
            {
                mark: [record for page in mark_pages for record in page]
                for mark, mark_pages in pages.items()
            },
            details,
        )
        self.pages = pages
        self.visited_pages: list[int] = []

    def search_pages(self, job: LookupJob, page_progress=None):
        for page_number, records in enumerate(self.pages[job.mark], start=1):
            self.visited_pages.append(page_number)
            if page_progress is not None:
                page_progress(page_number, len(self.pages[job.mark]))
            yield records


class FlakyClient(FakeClient):
    def __init__(
        self,
        records: dict[str, list[ResultRecord]],
        details: dict[int, DetailRecord],
        transient_mark: str,
        always_fail: bool = False,
    ) -> None:
        super().__init__(records, details)
        self.transient_mark = transient_mark
        self.always_fail = always_fail
        self.attempts: dict[str, int] = {}
        self.search_order: list[str] = []

    def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
        self.search_order.append(job.mark)
        self.attempts[job.mark] = self.attempts.get(job.mark, 0) + 1
        if job.mark == self.transient_mark and (
            self.always_fail or self.attempts[job.mark] == 1
        ):
            raise NeedsUserAction(BrowserState.WAITING_SITE, "结果页暂时不可用")
        return super().search(job, page_progress)


def test_application_skips_inactive_later_sequence_and_uses_active_registration(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [["云杉月付", "北京云杉网络技术有限公司", None, None]],
    )
    low = record(2, "云杉月付", "北京云杉网络技术有限公司", "66227440")
    high = record(4, "云杉月付", "北京云杉网络技术有限公司", "90000004")
    client = FakeClient(
        {"云杉月付": [low, high]},
        {4: detail(high, "2024年至2034年"), 2: active_detail(low, "2023年至2033年")},
    )
    before = sha256(source.read_bytes()).hexdigest()

    events: list[dict[str, object]] = []
    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("r1", source), events.append
    )

    assert client.opened_sequences == [4, 2]
    assert summary.state is TaskState.COMPLETED
    assert summary.items[0].status == "success"
    assert summary.items[0].term == "2023年至2033年"
    workbook = load_workbook(summary.output_path)
    assert workbook["Sheet1"]["C2"].value == "2023年至2033年"
    assert len(workbook["Sheet1"]._images) == 1
    workbook.close()
    assert sha256(source.read_bytes()).hexdigest() == before
    assert [
        event["stage"]
        for event in events
        if event.get("type") == "row_progress" and event.get("phase") == "stage"
    ] == [
        "searching",
        "checking_results",
        "opening_detail",
        "opening_detail",
        "capturing_evidence",
        "writing_workbook",
    ]


def test_application_stops_after_valid_candidate_on_first_page(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [["目标商标", "甲公司", None, None]],
    )
    lower = record(8, "目标商标", "甲公司")
    higher = record(10, "目标商标", "甲公司")
    later_page = record(20, "目标商标", "甲公司")
    client = PaginatedFakeClient(
        {"目标商标": [[lower, higher], [later_page]]},
        {
            10: active_detail(higher, "第一页有效期限"),
            20: active_detail(later_page, "第二页有效期限"),
        },
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("application-early-stop", source), lambda _: None
    )

    assert summary.items[0].status == "success"
    assert summary.items[0].term == "第一页有效期限"
    assert client.visited_pages == [1]
    assert client.opened_sequences == [10]


def test_application_continues_to_next_page_after_page_candidates_fail(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [["目标商标", "甲公司", None, None]],
    )
    inactive = record(9, "目标商标", "甲公司")
    active = record(7, "目标商标", "甲公司")
    never_needed = record(5, "目标商标", "甲公司")
    client = PaginatedFakeClient(
        {"目标商标": [[inactive], [active], [never_needed]]},
        {
            9: detail(inactive, None),
            7: active_detail(active, "第二页有效期限"),
            5: active_detail(never_needed, "第三页有效期限"),
        },
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("application-next-page", source), lambda _: None
    )

    assert summary.items[0].status == "success"
    assert summary.items[0].term == "第二页有效期限"
    assert client.visited_pages == [1, 2]
    assert client.opened_sequences == [9, 7]


def test_application_visits_every_page_before_returning_not_found(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [["目标商标", "甲公司", None, None]],
    )
    wrong_mark = record(11, "其他商标", "甲公司")
    wrong_applicant = record(10, "目标商标", "乙公司")
    client = PaginatedFakeClient(
        {"目标商标": [[wrong_mark], [wrong_applicant], []]},
        {},
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("application-all-pages", source), lambda _: None
    )

    assert summary.items[0].status == "not_found"
    assert summary.items[0].message == "未找到商标名称及申请人一致的检索记录"
    assert client.visited_pages == [1, 2, 3]
    assert client.opened_sequences == []


def test_application_not_found_reason_is_written_in_screenshot_column(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [["未找到商标", "原申请人", None, None]],
    )

    summary = TaskEngine(
        FakeClient({"未找到商标": []}), tmp_path / "data"
    ).run(TaskRequest("not-found-reason", source), lambda _: None)

    workbook = load_workbook(summary.output_path)
    sheet = workbook["Sheet1"]
    assert sheet["B2"].value == "原申请人"
    assert sheet["D2"].value == (
        "错误原因：未找到商标名称及申请人一致的检索记录"
    )
    assert len(sheet._images) == 0
    workbook.close()


def test_application_reports_inactive_status_separately(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [["无效商标", "甲公司", None, None]],
    )
    candidate = record(1, "无效商标", "甲公司")
    client = FakeClient(
        {"无效商标": [candidate]},
        {1: detail(candidate, "2020年至2030年")},
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("inactive-reason", source), lambda _: None
    )

    assert summary.items[0].message == "检索到申请人一致的商标，但状态不是有效注册商标"


def test_application_reports_missing_term_separately(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [["有效但无期限", "甲公司", None, None]],
    )
    candidate = record(1, "有效但无期限", "甲公司")
    client = FakeClient(
        {"有效但无期限": [candidate]},
        {1: active_detail(candidate, None)},
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("missing-term-reason", source), lambda _: None
    )

    assert summary.items[0].message == "检索到有效注册商标，但专用权期限为空"


def test_uniqueness_two_active_applicants_is_not_unique_and_uses_representatives(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.UNIQUENESS,
        [["青禾", None, "旧申请人", "旧期限", None]],
    )
    exact = [record(1, "青禾", "甲公司"), record(2, "青禾", "乙公司")]
    client = FakeClient(
        {"青禾": [record(9, "青禾月付", "甲公司"), *exact]},
        {1: active_detail(exact[0]), 2: active_detail(exact[1])},
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("r2", source), lambda _: None
    )

    result = summary.items[0]
    assert result.uniqueness == "否"
    assert result.applicant is None
    assert result.term is None
    assert client.opened_sequences == [2, 1]
    assert client.stacked_sequences == [2, 1]
    workbook = load_workbook(summary.output_path)
    sheet = workbook["Sheet1"]
    assert sheet["B2"].value == "否"
    assert sheet["C2"].value is None
    assert sheet["D2"].value is None
    assert len(sheet._images) == 1
    workbook.close()


def test_unique_company_fills_detail_from_first_active_candidate(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path, LookupMode.UNIQUENESS, [["青禾", None, None, None, None]]
    )
    low = record(3, "青禾", "北京青禾科技有限公司")
    high = record(8, "青禾", " 北京青禾科技有限公司 ")
    client = FakeClient(
        {"青禾": [low, high]},
        {8: detail(high, None), 3: active_detail(low)},
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("r3", source), lambda _: None
    )

    result = summary.items[0]
    assert result.uniqueness == "是"
    assert result.applicant == "北京青禾科技有限公司"
    assert result.term == "2025年至2035年"
    assert client.opened_sequences == [8, 3]


def test_uniqueness_multiple_applicants_only_one_active_is_unique(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path, LookupMode.UNIQUENESS, [["清泉饮品", None, None, None, None]]
    )
    owner_new = record(9, "清泉饮品", "山西示例水处理设备有限公司")
    owner_old = record(7, "清泉饮品", "山西示例水处理设备有限公司")
    other_new = record(8, "清泉饮品", "北京示例云科技有限公司")
    other_old = record(6, "清泉饮品", "北京示例云科技有限公司")
    client = FakeClient(
        {"清泉饮品": [owner_old, other_old, owner_new, other_new]},
        {
            9: active_detail(owner_new),
            8: detail(other_new, None),
            6: detail(other_old, None),
        },
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("r-mixed", source), lambda _: None
    )

    result = summary.items[0]
    assert result.uniqueness == "是"
    assert result.applicant == "山西示例水处理设备有限公司"
    assert client.opened_sequences == [9, 8, 6]
    assert all(sequence in {9, 7} for sequence in client.stacked_sequences)
    assert 1 <= len(client.stacked_sequences) <= 4


def test_uniqueness_stops_after_second_active_applicant(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path, LookupMode.UNIQUENESS, [["青松森林", None, None, None, None]]
    )
    first = record(10, "青松森林", "甲公司")
    second = record(9, "青松森林", "乙公司")
    never_needed = record(8, "青松森林", "丙公司")
    client = FakeClient(
        {"青松森林": [first, second, never_needed]},
        {10: active_detail(first), 9: active_detail(second), 8: active_detail(never_needed)},
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("r-short", source), lambda _: None
    )

    assert summary.items[0].uniqueness == "否"
    assert client.opened_sequences == [10, 9]
    assert set(client.stacked_sequences) == {10, 9}


def test_uniqueness_without_any_active_registration_needs_manual_review(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path, LookupMode.UNIQUENESS, [["无有效商标", None, None, None, None]]
    )
    first = record(2, "无有效商标", "甲公司")
    second = record(1, "无有效商标", "乙公司")
    client = FakeClient(
        {"无有效商标": [first, second]},
        {2: detail(first, None), 1: detail(second, None)},
    )

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("r-none", source), lambda _: None
    )

    result = summary.items[0]
    assert result.uniqueness == "未检索到有效注册商标，需人工复核"
    assert result.applicant is None
    assert result.term is None
    assert client.opened_sequences == [2, 1]


def test_only_user_requested_pause_stops_at_next_safe_row_boundary(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [
            ["商标一", "甲公司", None, None],
            ["商标二", "乙公司", None, None],
        ],
    )
    data_root = tmp_path / "data"
    first = record(1, "商标一", "甲公司")
    second = record(2, "商标二", "乙公司")
    entered = Event()
    release = Event()

    class BlockingClient(FakeClient):
        def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
            if job.mark == "商标一":
                entered.set()
                assert release.wait(timeout=3)
            return super().search(job, page_progress)

    engine = TaskEngine(
        BlockingClient(
            {"商标一": [first], "商标二": [second]},
            {1: active_detail(first, "第一项期限"), 2: active_detail(second, "第二项期限")},
        ),
        data_root,
    )
    events: list[dict[str, object]] = []
    summaries: list[object] = []

    def sink(event: dict[str, object]) -> None:
        if event["type"] == "paused":
            checkpoint = CheckpointStore(
                data_root / "tasks" / "pause" / "checkpoint.json"
            ).load()
            assert checkpoint is not None
            assert checkpoint.state == "paused_user_action"
            assert checkpoint.paused_row == 3
        events.append(event)

    thread = Thread(
        target=lambda: summaries.append(
            engine.run(TaskRequest("pause", source), sink)
        )
    )
    thread.start()
    assert entered.wait(timeout=3)
    engine.pause()
    release.set()
    thread.join(timeout=5)

    summary = summaries[0]
    assert summary.state is TaskState.PAUSED_USER_ACTION
    assert events[-1]["type"] == "paused"
    assert events[-1]["reason"] == "USER_REQUESTED"
    assert summary.completed_rows == (2,)


def test_login_required_pauses_immediately_without_querying_later_rows(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [
            ["商标一", "甲公司", None, None],
            ["需要登录", "乙公司", None, None],
            ["不应查询", "丙公司", None, None],
        ],
    )
    first = record(1, "商标一", "甲公司")
    searched: list[str] = []

    class LoginExpiryClient(FakeClient):
        def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
            searched.append(job.mark)
            return super().search(job, page_progress)

    data_root = tmp_path / "data"
    events: list[dict[str, object]] = []
    summary = TaskEngine(
        LoginExpiryClient(
            {"商标一": [first], "需要登录": [], "不应查询": []},
            {1: active_detail(first, "第一项期限")},
            pause_on_mark="需要登录",
        ),
        data_root,
    ).run(TaskRequest("login-expired", source), events.append)

    assert summary.state is TaskState.PAUSED_USER_ACTION
    assert summary.completed_rows == (2,)
    assert searched == ["商标一", "需要登录"]
    assert not any(event.get("type") == "row_deferred" for event in events)
    paused_event = events[-1]
    assert paused_event["type"] == "paused"
    assert paused_event["reason"] == "LOGIN_REQUIRED"
    assert paused_event["row"] == 3
    assert paused_event["output_path"] == str(summary.output_path)
    checkpoint = CheckpointStore(
        data_root / "tasks" / "login-expired" / "checkpoint.json"
    ).load()
    assert checkpoint is not None
    assert checkpoint.state == "paused_user_action"
    assert checkpoint.paused_row == 3


def test_row_error_is_deferred_then_retried_after_remaining_rows(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [
            ["商标一", "甲公司", None, None],
            ["商标二", "乙公司", None, None],
            ["商标三", "丙公司", None, None],
        ],
    )
    first = record(1, "商标一", "甲公司")
    second = record(2, "商标二", "乙公司")
    third = record(3, "商标三", "丙公司")
    client = FlakyClient(
        {"商标一": [first], "商标二": [second], "商标三": [third]},
        {
            1: active_detail(first, "第一项期限"),
            2: active_detail(second, "第二项期限"),
            3: active_detail(third, "第三项期限"),
        },
        transient_mark="商标二",
    )
    events: list[dict[str, object]] = []

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("defer-and-retry", source), events.append
    )

    assert summary.state is TaskState.COMPLETED
    assert client.search_order == ["商标一", "商标二", "商标三", "商标二"]
    assert [item.row for item in summary.items] == [2, 4, 3]
    assert all(item.status == "success" for item in summary.items)
    assert next(item for item in summary.items if item.row == 3).attempts == 2
    assert not any(event["type"] == "paused" for event in events)
    deferred = next(event for event in events if event["type"] == "row_deferred")
    assert (deferred["row"], deferred["attempt"], deferred["will_retry"]) == (
        3,
        1,
        True,
    )
    workbook = load_workbook(summary.output_path)
    log = workbook["查询执行记录"]
    assert [log.cell(row, 3).value for row in range(2, 5)] == [
        "成功",
        "成功",
        "成功",
    ]
    assert log.cell(3, 4).value == 2
    workbook.close()


def test_final_row_error_is_recorded_without_stopping_batch(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [
            ["始终失败", "甲公司", None, None],
            ["可以成功", "乙公司", None, None],
        ],
    )
    failed = record(1, "始终失败", "甲公司")
    success = record(2, "可以成功", "乙公司")
    client = FlakyClient(
        {"始终失败": [failed], "可以成功": [success]},
        {
            1: active_detail(failed, "失败项期限"),
            2: active_detail(success, "成功项期限"),
        },
        transient_mark="始终失败",
        always_fail=True,
    )
    events: list[dict[str, object]] = []

    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("record-final-failure", source), events.append
    )

    assert summary.state is TaskState.COMPLETED
    assert client.search_order == ["始终失败", "可以成功", "始终失败"]
    failed_item = next(item for item in summary.items if item.mark == "始终失败")
    assert failed_item.status == "failed"
    assert failed_item.attempts == 2
    assert "结果页暂时不可用" in failed_item.message
    assert not any(event["type"] == "paused" for event in events)
    workbook = load_workbook(summary.output_path)
    log = workbook["查询执行记录"]
    assert log.cell(2, 3).value == "失败"
    assert log.cell(2, 4).value == 2
    assert "结果页暂时不可用" in log.cell(2, 5).value
    assert workbook["Sheet1"]["D2"].value == "错误原因：结果页暂时不可用"
    assert workbook["Sheet1"]["C3"].value == "成功项期限"
    workbook.close()


def test_output_marks_rows_that_called_remote_model_even_after_retry_or_failure(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [
            ["无需模型", "甲公司", None, None],
            ["模型后重试", "乙公司", None, None],
            ["模型识别失败", "丙公司", None, None],
        ],
    )
    ordinary = record(1, "无需模型", "甲公司")
    recovered = record(2, "模型后重试", "乙公司")

    class MeteredClient(FakeClient):
        remote_model_call_count = 0
        attempts: dict[str, int] = {}

        def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
            self.attempts[job.mark] = self.attempts.get(job.mark, 0) + 1
            if job.mark == "模型后重试" and self.attempts[job.mark] == 1:
                self.remote_model_call_count += 1
                raise RuntimeError("识别后页面暂时不可用")
            if job.mark == "模型识别失败":
                self.remote_model_call_count += 1
                raise RuntimeError("模型识别失败")
            return super().search(job, page_progress)

    client = MeteredClient(
        {"无需模型": [ordinary], "模型后重试": [recovered]},
        {1: active_detail(ordinary), 2: active_detail(recovered)},
    )
    summary = TaskEngine(client, tmp_path / "data").run(
        TaskRequest("model-usage", source), lambda _event: None
    )

    workbook = load_workbook(summary.output_path)
    sheet = workbook["Sheet1"]
    assert sheet["F1"].value == "是否调用大模型"
    assert [sheet[f"F{row}"].value for row in (2, 3, 4)] == [
        "否", "是", "是"
    ]
    assert sheet["E1"].value == "查询结果"
    assert [sheet[f"E{row}"].value for row in (2, 3, 4)] == [
        "成功", "成功", None
    ]
    assert sheet.auto_filter.ref == "A1:F4"
    workbook.close()


def test_resume_skips_completed_rows_and_finishes_paused_row(tmp_path: Path) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [
            ["商标一", "甲公司", None, None],
            ["商标二", "乙公司", None, None],
        ],
    )
    first = record(1, "商标一", "甲公司")
    second = record(1, "商标二", "乙公司")
    data_root = tmp_path / "data"
    entered = Event()
    release = Event()

    class BlockingClient(FakeClient):
        def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
            if job.mark == "商标一" and not entered.is_set():
                entered.set()
                assert release.wait(timeout=3)
            return super().search(job, page_progress)

    client = BlockingClient(
        {"商标一": [first], "商标二": [second]},
        {1: active_detail(first, "2024年至2034年")},
    )
    engine = TaskEngine(client, data_root)
    paused_summaries: list[object] = []
    thread = Thread(
        target=lambda: paused_summaries.append(
            engine.run(TaskRequest("resume", source), lambda _: None)
        )
    )
    thread.start()
    assert entered.wait(timeout=3)
    engine.pause()
    release.set()
    thread.join(timeout=5)
    paused = paused_summaries[0]
    assert paused.state is TaskState.PAUSED_USER_ACTION

    client.details[1] = active_detail(second, "2025年至2035年")
    client.opened_sequences.clear()
    events: list[dict[str, object]] = []
    completed = engine.resume("resume", events.append)

    assert completed.state is TaskState.COMPLETED
    assert len(completed.items) == 1
    assert completed.items[0].row == 3
    workbook = load_workbook(completed.output_path)
    assert workbook["Sheet1"]["C2"].value == "2024年至2034年"
    assert workbook["Sheet1"]["C3"].value == "2025年至2035年"
    workbook.close()
    resumed_row = next(
        event
        for event in events
        if event.get("type") == "row_progress" and event.get("phase") == "started"
    )
    assert (resumed_row["current"], resumed_row["total"]) == (2, 2)


def test_twelve_application_rows_finish_in_order_with_one_embedded_image_each(
    tmp_path: Path,
) -> None:
    rows = [[f"商标{index}", f"第{index}家公司", None, None] for index in range(1, 13)]
    source = make_workbook(tmp_path, LookupMode.APPLICATION, rows)
    records: dict[str, list[ResultRecord]] = {}
    details: dict[int, DetailRecord] = {}
    for index in range(1, 13):
        item = record(index, f"商标{index}", f"第{index}家公司")
        records[item.mark] = [item]
        details[index] = active_detail(item, f"第{index}项期限")

    events: list[dict[str, object]] = []
    summary = TaskEngine(FakeClient(records, details), tmp_path / "data").run(
        TaskRequest("batch-12", source), events.append
    )

    assert summary.state is TaskState.COMPLETED
    assert [item.row for item in summary.items] == list(range(2, 14))
    completed = [
        event
        for event in events
        if event.get("type") == "row_progress" and event.get("phase") == "completed"
    ]
    assert [(event["current"], event["total"]) for event in completed] == [
        (index, 12) for index in range(1, 13)
    ]
    workbook = load_workbook(summary.output_path)
    sheet = workbook["Sheet1"]
    assert [sheet.cell(row, 3).value for row in range(2, 14)] == [
        f"第{index}项期限" for index in range(1, 13)
    ]
    assert len(sheet._images) == 12
    assert [image.anchor._from.row for image in sheet._images] == list(range(1, 13))
    assert all((sheet.row_dimensions[row].height or 0) <= 409 for row in range(2, 14))
    workbook.close()


def test_twelve_uniqueness_rows_finish_in_order_with_stacked_evidence(
    tmp_path: Path,
) -> None:
    rows = [[f"唯一商标{index}", None, None, None, None] for index in range(1, 13)]
    source = make_workbook(tmp_path, LookupMode.UNIQUENESS, rows)
    records: dict[str, list[ResultRecord]] = {}
    details: dict[int, DetailRecord] = {}
    for index in range(1, 13):
        item = record(index, f"唯一商标{index}", f"第{index}家公司")
        records[item.mark] = [item]
        details[index] = active_detail(item, f"第{index}项期限")

    events: list[dict[str, object]] = []
    summary = TaskEngine(FakeClient(records, details), tmp_path / "data").run(
        TaskRequest("uniqueness-batch-12", source), events.append
    )

    assert summary.state is TaskState.COMPLETED
    assert [item.row for item in summary.items] == list(range(2, 14))
    workbook = load_workbook(summary.output_path)
    sheet = workbook["Sheet1"]
    assert [sheet.cell(row, 2).value for row in range(2, 14)] == ["是"] * 12
    assert [sheet.cell(row, 3).value for row in range(2, 14)] == [
        f"第{index}家公司" for index in range(1, 13)
    ]
    assert [sheet.cell(row, 4).value for row in range(2, 14)] == [
        f"第{index}项期限" for index in range(1, 13)
    ]
    assert len(sheet._images) == 12
    assert [image.anchor._from.row for image in sheet._images] == list(range(1, 13))
    workbook.close()


def test_cancel_emits_partial_result_location_after_current_row_finishes(
    tmp_path: Path,
) -> None:
    source = make_workbook(
        tmp_path,
        LookupMode.APPLICATION,
        [
            ["商标一", "甲公司", None, None],
            ["商标二", "乙公司", None, None],
        ],
    )
    first = record(1, "商标一", "甲公司")
    second = record(2, "商标二", "乙公司")
    entered = Event()
    release = Event()

    class BlockingClient(FakeClient):
        def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
            if job.mark == "商标一":
                entered.set()
                assert release.wait(timeout=3)
            return super().search(job, page_progress)

    engine = TaskEngine(
        BlockingClient(
            {"商标一": [first], "商标二": [second]},
            {1: active_detail(first, "第一项期限"), 2: active_detail(second, "第二项期限")},
        ),
        tmp_path / "data",
    )
    events: list[dict[str, object]] = []
    summaries: list[object] = []
    thread = Thread(
        target=lambda: summaries.append(
            engine.run(TaskRequest("cancel-batch", source), events.append)
        )
    )
    thread.start()
    assert entered.wait(timeout=3)
    engine.cancel()
    release.set()
    thread.join(timeout=5)

    assert not thread.is_alive()
    summary = summaries[0]
    assert summary.state is TaskState.CANCELLED
    assert summary.completed_rows == (2,)
    cancelled_event = next(
        event
        for event in events
        if event.get("type") == "task_state" and event.get("state") == "cancelled"
    )
    assert (cancelled_event["current"], cancelled_event["total"]) == (1, 2)
    assert cancelled_event["output_path"] == str(summary.output_path)
