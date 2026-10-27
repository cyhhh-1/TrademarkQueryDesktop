from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor
from PIL import Image

from trademark_worker.models import LookupJob, RowResult
from trademark_worker.workbook.inspector import inspect_workbook, prepare_working_copy
from trademark_worker.workbook.writer import (
    WorkbookOutputError,
    apply_row_result,
    record_execution_status,
    validate_output,
    write_row_error,
)


def make_workbook(tmp_path: Path, mode: str) -> tuple[Path, object, LookupJob]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    if mode == "application":
        sheet.append(["商标名称", "商标申请人", "专用权期限", "截图"])
        sheet.append(["云杉月付", "输入申请人", None, None])
        job = LookupJob.application(2, "云杉月付", "输入申请人")
        name = "商标申请.xlsx"
    else:
        sheet.append(
            [
                "商标名称",
                "是否唯一性\n（当申请公司有多家时填写“否”",
                "商标申请人\n当是否唯一性“否”时此处空白",
                "商标有效期",
                "截图",
            ]
        )
        sheet.append(["青禾月付", None, None, None, None])
        job = LookupJob.uniqueness(2, "青禾月付")
        name = "商标查询是否唯一性.xlsx"
    path = tmp_path / name
    workbook.save(path)
    return path, inspect_workbook(path), job


def evidence_png(tmp_path: Path, size: tuple[int, int] = (640, 900)) -> Path:
    path = tmp_path / "evidence.png"
    Image.new("RGB", size, "white").save(path)
    return path


def test_application_result_writes_term_and_embeds_detail_image(
    tmp_path: Path,
) -> None:
    source, inspection, job = make_workbook(tmp_path, "application")
    original_hash = sha256(source.read_bytes()).hexdigest()
    working = prepare_working_copy(source, tmp_path / "run")

    apply_row_result(
        working,
        inspection,
        job,
        RowResult.application_success(
            "北京云杉网络技术有限公司",
            "2023年08月28日至2033年08月27日",
            "66227440",
        ),
        evidence_png(tmp_path),
    )

    workbook = load_workbook(working)
    sheet = workbook["Sheet1"]
    assert sheet["B2"].value == "北京云杉网络技术有限公司"
    assert sheet["C2"].value == "2023年08月28日至2033年08月27日"
    assert sheet["E1"].value == "查询结果"
    assert sheet["E2"].value == "成功"
    assert len(sheet._images) == 1
    assert isinstance(sheet._images[0].anchor, TwoCellAnchor)
    assert sheet._images[0].anchor.editAs == "twoCell"
    assert sheet._images[0].anchor._from.row == 1
    assert sheet._images[0].anchor._from.col == 3
    assert sheet._images[0].anchor.to.row == 2
    assert sheet._images[0].anchor.to.col == 4
    assert 300 <= sheet.row_dimensions[2].height <= 409
    assert sheet.column_dimensions["B"].width >= 28
    assert sheet.column_dimensions["C"].width >= 32
    assert sheet.column_dimensions["D"].width >= 90
    workbook.close()
    assert sha256(source.read_bytes()).hexdigest() == original_hash


def test_application_failure_adds_result_header_but_leaves_value_blank(
    tmp_path: Path,
) -> None:
    source, inspection, job = make_workbook(tmp_path, "application")
    working = prepare_working_copy(source, tmp_path / "run")

    write_row_error(working, inspection, job, "未找到有效商标")

    workbook = load_workbook(working)
    sheet = workbook["Sheet1"]
    assert sheet["D2"].value == "错误原因：未找到有效商标"
    assert sheet["E1"].value == "查询结果"
    assert sheet["E2"].value is None
    workbook.close()


def test_non_unique_result_leaves_applicant_and_term_blank(tmp_path: Path) -> None:
    source, inspection, job = make_workbook(tmp_path, "uniqueness")
    working = prepare_working_copy(source, tmp_path / "run")

    apply_row_result(
        working,
        inspection,
        job,
        RowResult.uniqueness_result("否", applicant="错误公司", term="错误期限"),
        evidence_png(tmp_path, (800, 120)),
    )
    record_execution_status(
        working, job, "success", 1,
        inspection=inspection, model_used=False,
    )

    workbook = load_workbook(working)
    sheet = workbook["Sheet1"]
    assert sheet["B2"].value == "否"
    assert sheet["C2"].value is None
    assert sheet["D2"].value is None
    assert sheet["F1"].value == "是否调用大模型"
    assert sheet["F2"].value == "否"
    assert sheet.auto_filter.ref == "A1:F2"
    assert len(sheet._images) == 1
    workbook.close()


def test_tall_evidence_keeps_full_resolution_inside_xlsx(tmp_path: Path) -> None:
    source, inspection, job = make_workbook(tmp_path, "uniqueness")
    working = prepare_working_copy(source, tmp_path / "run")
    evidence = evidence_png(tmp_path, (640, 5000))

    apply_row_result(
        working,
        inspection,
        job,
        RowResult.uniqueness_result("否"),
        evidence,
    )

    with ZipFile(working) as archive:
        media = next(name for name in archive.namelist() if name.startswith("xl/media/"))
        with archive.open(media) as stream, Image.open(stream) as embedded:
            assert embedded.size == (640, 5000)

    workbook = load_workbook(working)
    sheet = workbook["Sheet1"]
    assert sheet.row_dimensions[2].height <= 409
    assert sheet.column_dimensions["E"].width >= 90
    workbook.close()


def test_unique_result_may_write_applicant_and_term(tmp_path: Path) -> None:
    source, inspection, job = make_workbook(tmp_path, "uniqueness")
    working = prepare_working_copy(source, tmp_path / "run")

    apply_row_result(
        working,
        inspection,
        job,
        RowResult.uniqueness_result(
            "是", "北京青禾科技有限公司", "2025年至2035年", "90000001"
        ),
        evidence_png(tmp_path, (800, 120)),
    )

    workbook = load_workbook(working)
    sheet = workbook["Sheet1"]
    assert [sheet.cell(2, column).value for column in range(2, 5)] == [
        "是",
        "北京青禾科技有限公司",
        "2025年至2035年",
    ]
    assert sheet.column_dimensions["C"].width >= 28
    assert sheet.column_dimensions["D"].width >= 32
    workbook.close()


def test_validate_output_rejects_non_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    path.write_text("not a zip", encoding="utf-8")

    with pytest.raises(WorkbookOutputError, match="有效的 XLSX"):
        validate_output(path)


def test_missing_evidence_is_rejected_before_workbook_change(tmp_path: Path) -> None:
    source, inspection, job = make_workbook(tmp_path, "application")
    working = prepare_working_copy(source, tmp_path / "run")
    before = sha256(working.read_bytes()).hexdigest()

    with pytest.raises(WorkbookOutputError, match="证据图片不存在"):
        apply_row_result(
            working,
            inspection,
            job,
            RowResult.application_success("申请人", "期限"),
            tmp_path / "missing.png",
        )

    assert sha256(working.read_bytes()).hexdigest() == before
