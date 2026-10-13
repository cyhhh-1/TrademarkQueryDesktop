from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image

from trademark_worker.models import DetailRecord, LookupJob, ResultRecord, TaskRequest
from trademark_worker.task_engine import TaskEngine


class OfflineClient:
    def __init__(self) -> None:
        self.application = ResultRecord(
            1,
            "66227440",
            "36",
            "2022年07月27日",
            "抖音月付",
            "北京字跳网络技术有限公司",
            row_index=0,
        )
        self.uniqueness = [
            ResultRecord(
                1,
                "93296493",
                "35",
                "2026年08月03日",
                "美团",
                "北京三快科技有限公司",
                row_index=0,
            ),
            ResultRecord(
                2,
                "93296492",
                "25",
                "2026年08月03日",
                "美团",
                "北京三快科技有限公司",
                row_index=1,
            ),
        ]

    def search(self, job: LookupJob, page_progress=None) -> list[ResultRecord]:
        return [self.application] if job.mark == "抖音月付" else self.uniqueness

    def search_pages(self, job: LookupJob, page_progress=None):
        yield self.search(job, page_progress)

    def open_detail(self, record: ResultRecord) -> DetailRecord:
        term = (
            "2023年08月28日至2033年08月27日"
            if record.mark == "抖音月付"
            else "2026年至2036年"
        )
        status = "LIVE/REGISTRATION/Issued and Active 注册商标"
        return DetailRecord(
            record.mark, record.registration_no, record.applicant, term, status=status
        )

    def capture_detail(self, output: Path) -> Path:
        output.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (640, 900), "white").save(output)
        return output

    def capture_exact_rows(
        self, records: list[ResultRecord], output_dir: Path
    ) -> Path:
        output = output_dir / "exact-results.png"
        output.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (800, 120 * len(records)), "white").save(output)
        return output


def application_workbook(tmp_path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["商标名称", "商标申请人", "专用权期限", "截图"])
    sheet.append(["抖音月付", "北京字跳网络技术有限公司", None, None])
    path = tmp_path / "商标申请.xlsx"
    workbook.save(path)
    return path


def uniqueness_workbook(tmp_path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "商标名称",
            "是否唯一性\n（当申请公司有多家时填写“否”",
            "商标申请人\n当是否唯一性“否”时此处空白",
            "商标有效期",
            "截图",
        ]
    )
    sheet.append(["美团", None, None, None, None])
    path = tmp_path / "商标查询是否唯一性.xlsx"
    workbook.save(path)
    return path


def test_offline_application_finishes_without_touching_source(tmp_path: Path) -> None:
    source = application_workbook(tmp_path)
    before = sha256(source.read_bytes()).hexdigest()

    summary = TaskEngine(OfflineClient(), tmp_path / "data").run(
        TaskRequest("application", source), lambda _: None
    )

    workbook = load_workbook(summary.output_path)
    assert workbook.active["C2"].value == "2023年08月28日至2033年08月27日"
    assert len(workbook.active._images) == 1
    workbook.close()
    assert sha256(source.read_bytes()).hexdigest() == before


def test_offline_uniqueness_writes_decision_and_stacked_evidence(tmp_path: Path) -> None:
    source = uniqueness_workbook(tmp_path)

    summary = TaskEngine(OfflineClient(), tmp_path / "data").run(
        TaskRequest("uniqueness", source), lambda _: None
    )

    workbook = load_workbook(summary.output_path)
    sheet = workbook.active
    assert sheet["B2"].value == "是"
    assert sheet["C2"].value == "北京三快科技有限公司"
    assert sheet["D2"].value == "2026年至2036年"
    assert len(sheet._images) == 1
    workbook.close()
